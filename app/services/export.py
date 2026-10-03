"""ATS-safe resume export: plain single-column DOCX and PDF, built directly from
the canonical ResumeData - no tables, no text boxes, no columns, no headers/footers,
no embedded images. Spacing is kept tight so a typical resume fits one page,
matching standard ATS-template conventions (Calibri/Helvetica, 10-11pt body,
single-column, standard section headers).
"""
import io
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.schemas.resume import ResumeData

ACCENT_HEX = "#4F46E5"  # indigo-600
ACCENT_RGB = RGBColor(0x4F, 0x46, 0xE5)
FONT_NAME = "Calibri"


def _esc(text: str | None) -> str:
    """reportlab's Paragraph treats text as mini-XML - escape & < > so raw
    characters in resume content (e.g. 'R&D') don't break rendering."""
    return escape(text or "")


def _contact_line(data: ResumeData) -> str:
    parts = [p for p in [data.contact.email, data.contact.phone, data.contact.location] if p]
    links = [p for p in [data.contact.linkedin, data.contact.website] if p]
    return "  |  ".join(parts + links)


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def generate_ats_docx(data: ResumeData) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)
        section.left_margin = Pt(44)
        section.right_margin = Pt(44)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.line_spacing = 1.08

    def heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = ACCENT_RGB
        pPr = p._p.get_or_add_pPr()
        pBdr = pPr.makeelement(qn("w:pBdr"), {})
        bottom = pPr.makeelement(qn("w:bottom"), {
            qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "1", qn("w:color"): "4F46E5",
        })
        pBdr.append(bottom)
        pPr.append(pBdr)

    def bullet(text: str):
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05

    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_after = Pt(0)
    name_run = name_p.add_run(data.name or "")
    name_run.bold = True
    name_run.font.size = Pt(19)
    name_run.font.color.rgb = ACCENT_RGB

    if data.title:
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_after = Pt(2)
        title_run = title_p.add_run(data.title)
        title_run.font.size = Pt(11.5)
        title_run.italic = True

    contact = _contact_line(data)
    if contact:
        c = doc.add_paragraph(contact)
        c.paragraph_format.space_after = Pt(4)
        c.runs[0].font.size = Pt(9.5)

    if data.summary:
        heading("Summary")
        doc.add_paragraph(data.summary)

    if data.skills:
        heading("Skills")
        doc.add_paragraph(", ".join(data.skills))

    if data.experiences:
        heading("Experience")
        for exp in data.experiences:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(1)
            header = " — ".join([v for v in (exp.role, exp.company) if v])
            run = p.add_run(header)
            run.bold = True
            run.font.size = Pt(10.5)
            dates = " – ".join([v for v in (exp.start_date, exp.end_date) if v])
            if dates:
                date_run = p.add_run(f"    {dates}")
                date_run.italic = True
                date_run.font.size = Pt(9)
            for b in exp.bullets:
                bullet(b)

    if data.projects:
        heading("Projects")
        for proj in data.projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(proj.name)
            run.bold = True
            if proj.link:
                link_run = p.add_run(f"  —  {proj.link}")
                link_run.font.size = Pt(9)
            if proj.description:
                doc.add_paragraph(proj.description)
            if proj.tech_stack:
                tech_p = doc.add_paragraph(", ".join(proj.tech_stack))
                tech_p.runs[0].italic = True
                tech_p.runs[0].font.size = Pt(9)

    if data.education:
        heading("Education")
        for ed in data.education:
            line = ed.degree
            if ed.institution:
                line += f", {ed.institution}"
            if ed.year:
                line += f" ({ed.year})"
            doc.add_paragraph(line)

    if data.certificates:
        heading("Certificates")
        for c in data.certificates:
            line = c.name
            if c.issuer:
                line += f", {c.issuer}"
            if c.year:
                line += f" ({c.year})"
            doc.add_paragraph(line)

    if data.achievements:
        heading("Achievements")
        for a in data.achievements:
            line = a.description
            if a.date:
                line += f" ({a.date})"
            doc.add_paragraph(line)

    if data.languages:
        heading("Languages")
        doc.add_paragraph(", ".join(data.languages))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF (reportlab - pure Python, no system binary needed, safe on Render)
# --------------------------------------------------------------------------- #
def generate_ats_pdf(data: ResumeData) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=0.55 * inch,
        rightMargin=0.55 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
    )

    accent = colors.HexColor(ACCENT_HEX)
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "Name", parent=styles["Title"], alignment=TA_LEFT, fontName="Helvetica-Bold",
        fontSize=19, textColor=accent, spaceAfter=1, leading=22,
    )
    title_style = ParagraphStyle(
        "JobTitle", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=11.5, spaceAfter=3,
    )
    normal = ParagraphStyle("NormalATS", parent=styles["Normal"], fontName="Helvetica", fontSize=10, leading=12.5, spaceAfter=1)
    contact_style = ParagraphStyle("Contact", parent=normal, fontSize=9.5, spaceAfter=5)
    small_italic = ParagraphStyle("SmallItalic", parent=normal, fontName="Helvetica-Oblique", fontSize=9, textColor=colors.HexColor("#555555"))
    heading_style = ParagraphStyle(
        "Heading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11.5,
        textColor=accent, spaceBefore=8, spaceAfter=1,
    )
    bold_line = ParagraphStyle("BoldLine", parent=normal, fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=4)

    def heading(text: str):
        story.append(Paragraph(text.upper(), heading_style))
        story.append(HRFlowable(width="100%", thickness=0.75, color=accent, spaceAfter=3))

    story = [Paragraph(_esc(data.name), name_style)]
    if data.title:
        story.append(Paragraph(_esc(data.title), title_style))
    contact = _contact_line(data)
    if contact:
        story.append(Paragraph(_esc(contact), contact_style))

    if data.summary:
        heading("Summary")
        story.append(Paragraph(_esc(data.summary), normal))

    if data.skills:
        heading("Skills")
        story.append(Paragraph(_esc(", ".join(data.skills)), normal))

    if data.experiences:
        heading("Experience")
        for exp in data.experiences:
            header = " — ".join([v for v in (exp.role, exp.company) if v])
            dates = " – ".join([v for v in (exp.start_date, exp.end_date) if v])
            story.append(Paragraph(_esc(header), bold_line))
            if dates:
                story.append(Paragraph(_esc(dates), small_italic))
            if exp.bullets:
                items = [ListItem(Paragraph(_esc(b), normal), spaceBefore=0, spaceAfter=0) for b in exp.bullets]
                story.append(ListFlowable(items, bulletType="bullet", leftIndent=14, bulletColor=accent, spaceBefore=1, spaceAfter=1))
            story.append(Spacer(1, 2))

    if data.projects:
        heading("Projects")
        for proj in data.projects:
            line = proj.name + (f"  —  {proj.link}" if proj.link else "")
            story.append(Paragraph(_esc(line), bold_line))
            if proj.description:
                story.append(Paragraph(_esc(proj.description), normal))
            if proj.tech_stack:
                story.append(Paragraph(_esc(", ".join(proj.tech_stack)), small_italic))
            story.append(Spacer(1, 2))

    if data.education:
        heading("Education")
        for ed in data.education:
            line = ed.degree
            if ed.institution:
                line += f", {ed.institution}"
            if ed.year:
                line += f" ({ed.year})"
            story.append(Paragraph(_esc(line), normal))

    if data.certificates:
        heading("Certificates")
        for c in data.certificates:
            line = c.name
            if c.issuer:
                line += f", {c.issuer}"
            if c.year:
                line += f" ({c.year})"
            story.append(Paragraph(_esc(line), normal))

    if data.achievements:
        heading("Achievements")
        for a in data.achievements:
            line = a.description + (f" ({a.date})" if a.date else "")
            story.append(Paragraph(_esc(line), normal))

    if data.languages:
        heading("Languages")
        story.append(Paragraph(_esc(", ".join(data.languages)), normal))

    doc.build(story)
    return buf.getvalue()